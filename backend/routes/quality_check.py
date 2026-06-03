from io import BytesIO
from datetime import datetime
from flask import Blueprint, request, jsonify, send_file
from flask_jwt_extended import jwt_required
from sqlalchemy.orm import joinedload
from models import db, Task, Unit, QualityIssue, ConfirmedPattern, ProxyMetricPair
from sqlalchemy import or_
from services.quality_checker import run_quality_check, _hash_text, get_proxy_issues
from services.permission import require_perm
try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None

qc_bp = Blueprint("quality_check", __name__)


def _issue_to_dict(i):
    unit_name = ""
    assessor_unit_name = ""
    main_task = ""
    if i.task:
        if i.task.unit:
            unit_name = i.task.unit.name
        if i.task.assessor_unit:
            assessor_unit_name = i.task.assessor_unit.name
        main_task = i.task.main_task or ""
    return {
        "id": i.id, "plan_id": i.plan_id, "task_id": i.task_id,
        "issue_type": i.issue_type, "column_name": i.column_name,
        "text": i.text, "suggestion": i.suggestion, "context": i.context,
        "confidence": i.confidence, "status": i.status,
        "unit_name": unit_name,
        "assessor_unit_name": assessor_unit_name,
        "main_task": main_task,
    }


@qc_bp.route("/api/quality-check/run", methods=["POST"])
@jwt_required()
@require_perm("quality_manage")
def run_check():
    data = request.get_json() or {}
    plan_id = data.get("plan_id")
    if not plan_id:
        return jsonify({"msg": "请选择考核方案"}), 400
    results = run_quality_check(plan_id=plan_id)
    return jsonify({"data": results, "msg": f"检测完成，共发现 {results['total_issues']} 个问题"})


@qc_bp.route("/api/quality-check/issues", methods=["GET"])
@jwt_required()
def list_issues():
    plan_id = request.args.get("plan_id", type=int)
    status = request.args.get("status", "")
    issue_type = request.args.get("issue_type", "")
    assessor_unit_id = request.args.get("assessor_unit_id", type=int)
    unit_id = request.args.get("unit_id", type=int)
    search = request.args.get("search", "")
    page = request.args.get("page", 1, type=int)
    page_size = request.args.get("page_size", 100, type=int)
    page_size = min(page_size, 500)

    q = QualityIssue.query.options(
        joinedload(QualityIssue.task).joinedload(Task.unit),
        joinedload(QualityIssue.task).joinedload(Task.assessor_unit),
    )
    if plan_id:
        q = q.filter(QualityIssue.plan_id == plan_id)
    if status:
        q = q.filter(QualityIssue.status == status)
    if issue_type:
        q = q.filter(QualityIssue.issue_type == issue_type)
    if assessor_unit_id:
        q = q.join(QualityIssue.task).filter(Task.assessor_unit_id == assessor_unit_id)
    if unit_id:
        q = q.join(QualityIssue.task).filter(Task.unit_id == unit_id)
    if search:
        q = q.filter(
            or_(
                QualityIssue.text.contains(search),
                QualityIssue.suggestion.contains(search),
                QualityIssue.context.contains(search),
            )
        )

    total = q.count()
    items = q.order_by(QualityIssue.id).offset((page - 1) * page_size).limit(page_size).all()
    return jsonify({"data": {
        "items": [_issue_to_dict(i) for i in items],
        "total": total, "page": page, "page_size": page_size,
    }})


@qc_bp.route("/api/quality-check/proxy-issues", methods=["GET"])
@jwt_required()
def list_proxy_issues():
    """获取疑似以指标考指标问题（从 ProxyMetricPair 表读取）"""
    plan_id = request.args.get("plan_id", type=int)
    if not plan_id:
        return jsonify({"data": {"items": [], "total": 0}})
    items = get_proxy_issues(plan_id)
    return jsonify({"data": {"items": items, "total": len(items)}})


@qc_bp.route("/api/quality-check/issues/<int:issue_id>", methods=["PUT"])
@jwt_required()
@require_perm("quality_manage")
def update_issue(issue_id):
    i = QualityIssue.query.get(issue_id)
    if not i:
        return jsonify({"msg": "问题不存在"}), 404
    data = request.get_json() or {}
    if "status" in data:
        i.status = data["status"]
        if data["status"] == "confirmed" and i.text:
            h = _hash_text(i.text)
            if not ConfirmedPattern.query.filter_by(
                plan_id=i.plan_id, issue_type=i.issue_type, text_hash=h
            ).first():
                db.session.add(ConfirmedPattern(
                    plan_id=i.plan_id, issue_type=i.issue_type, text_hash=h
                ))
        if data["status"] == "pending" and i.text:
            h = _hash_text(i.text)
            ConfirmedPattern.query.filter_by(
                plan_id=i.plan_id, issue_type=i.issue_type, text_hash=h
            ).delete()
    db.session.commit()
    return jsonify({"msg": "更新成功"})


@qc_bp.route("/api/quality-check/issues/batch-status", methods=["POST"])
@jwt_required()
@require_perm("quality_manage")
def batch_update_status():
    data = request.get_json() or {}
    ids = data.get("ids", [])
    status = data.get("status", "")
    if not ids or not status:
        return jsonify({"msg": "参数错误"}), 400
    QualityIssue.query.filter(QualityIssue.id.in_(ids)).update({"status": status}, synchronize_session=False)
    db.session.commit()
    return jsonify({"msg": f"已更新 {len(ids)} 条问题"})


@qc_bp.route("/api/quality-check/batch-correct", methods=["POST"])
@jwt_required()
@require_perm("quality_manage")
def batch_correct():
    """更正任务字段，同步更新同 assessor + 同 main_task 的所有任务"""
    data = request.get_json() or {}
    task_id = data.get("task_id")
    field = data.get("field")
    new_value = data.get("new_value", "")
    updates = data.get("updates", {})
    if not task_id:
        return jsonify({"msg": "参数错误"}), 400

    task = Task.query.get(task_id)
    if not task:
        return jsonify({"msg": "任务不存在"}), 404

    all_updates = dict(updates)
    if field and field != "_noop":
        all_updates[field] = new_value
    is_noop = (field == "_noop" and not updates)

    plan_id = task.plan_id
    assessor_id = task.assessor_unit_id

    siblings_q = Task.query.filter(
        Task.plan_id == plan_id,
        Task.assessor_unit_id == assessor_id,
        Task.main_task == task.main_task,
    )

    if not is_noop and all_updates:
        for f, new_val in all_updates.items():
            if hasattr(Task, f):
                old_val = getattr(task, f, "")
                siblings_q = siblings_q.filter(getattr(Task, f) == old_val)

    siblings = siblings_q.all()

    text_fields = {"key_work", "main_task", "scoring_note"}
    updated = 0
    if not is_noop and all_updates:
        for t in siblings:
            for f, new_val in all_updates.items():
                if f in text_fields and hasattr(t, f):
                    setattr(t, f, new_val)
            updated += 1
        for f in ["assessment_dimension_id", "unit_id", "assessor_unit_id", "review_period"]:
            if f in all_updates and hasattr(task, f):
                setattr(task, f, all_updates[f])

    resolved = QualityIssue.query.filter(
        QualityIssue.task_id.in_([t.id for t in siblings]),
        QualityIssue.status == "pending",
    ).update({"status": "resolved"}, synchronize_session=False)

    db.session.commit()
    return jsonify({
        "msg": f"已更正 {updated} 条任务，解决 {resolved} 个问题，涉及 {len(set(t.unit_id for t in siblings))} 个单位",
        "resolved_count": resolved, "task_count": updated,
        "unit_count": len(set(t.unit_id for t in siblings)),
    })


@qc_bp.route("/api/quality-check/batch-confirm", methods=["POST"])
@jwt_required()
@require_perm("quality_manage")
def batch_confirm():
    """确认无误：将同一主考单位、同一主要任务下所有待处理问题标记为确认无误"""
    data = request.get_json() or {}
    task_id = data.get("task_id")
    if not task_id:
        return jsonify({"msg": "参数错误"}), 400

    task = Task.query.get(task_id)
    if not task:
        return jsonify({"msg": "任务不存在"}), 404

    siblings = Task.query.filter(
        Task.plan_id == task.plan_id,
        Task.assessor_unit_id == task.assessor_unit_id,
        Task.main_task == task.main_task,
    ).all()
    sibling_ids = [t.id for t in siblings]

    pending_issues = QualityIssue.query.filter(
        QualityIssue.task_id.in_(sibling_ids),
        QualityIssue.status == "pending",
    ).all()

    existing_cp = set(
        (p.issue_type, p.text_hash)
        for p in ConfirmedPattern.query.filter_by(plan_id=task.plan_id).all()
    )
    confirmed_count = 0
    for iss in pending_issues:
        iss.status = "confirmed"
        if iss.text:
            h = _hash_text(iss.text)
            key = (iss.issue_type, h)
            if key not in existing_cp:
                db.session.add(ConfirmedPattern(
                    plan_id=task.plan_id, issue_type=iss.issue_type, text_hash=h
                ))
                existing_cp.add(key)
        confirmed_count += 1

    db.session.commit()
    return jsonify({
        "msg": f"已确认 {confirmed_count} 条问题为误报，涉及 {len(siblings)} 条任务、{len(set(t.unit_id for t in siblings))} 个单位",
        "confirmed_count": confirmed_count, "task_count": len(siblings),
        "unit_count": len(set(t.unit_id for t in siblings)),
    })


@qc_bp.route("/api/quality-check/correct-single", methods=["POST"])
@jwt_required()
@require_perm("quality_manage")
def correct_single():
    """保存并更正当前：仅更正单条任务，标记其相关问题为已修复"""
    data = request.get_json() or {}
    task_id = data.get("task_id")
    updates = data.get("updates", {})
    if not task_id or not updates:
        return jsonify({"msg": "参数错误"}), 400

    task = Task.query.get(task_id)
    if not task:
        return jsonify({"msg": "任务不存在"}), 404

    allowed_fields = {
        "key_work", "main_task", "scoring_note",
        "assessment_dimension_id", "unit_id", "assessor_unit_id", "review_period",
    }
    for field, new_value in updates.items():
        if field in allowed_fields and hasattr(task, field):
            setattr(task, field, new_value)

    resolved = QualityIssue.query.filter(
        QualityIssue.task_id == task_id,
        QualityIssue.status == "pending",
    ).update({"status": "resolved"}, synchronize_session=False)

    db.session.commit()
    return jsonify({
        "msg": f"已更正单条任务，解决 {resolved} 个问题",
        "resolved_count": resolved,
    })


@qc_bp.route("/api/quality-check/proxy-metrics/confirm", methods=["POST"])
@jwt_required()
@require_perm("quality_manage")
def confirm_proxy_metric():
    """确认疑似以指标考指标 pair 为误报"""
    data = request.get_json() or {}
    plan_id = data.get("plan_id")
    task_id_a = data.get("task_id_a")
    task_id_b = data.get("task_id_b")
    if not plan_id or not task_id_a or not task_id_b:
        return jsonify({"msg": "参数错误"}), 400

    tid_a = min(task_id_a, task_id_b)
    tid_b = max(task_id_a, task_id_b)

    # 在 ProxyMetricPair 表中查找并确认
    pair = ProxyMetricPair.query.filter_by(
        plan_id=plan_id, task_id_a=tid_a, task_id_b=tid_b,
    ).first()
    if pair:
        pair.status = "confirmed"
        pair.updated_at = db.func.now()
    else:
        # 兼容：如果没有已有记录，创建一条 confirmed 记录
        db.session.add(ProxyMetricPair(
            plan_id=plan_id, task_id_a=tid_a, task_id_b=tid_b,
            status="confirmed", similarity=0, confidence="medium",
        ))
    db.session.commit()

    return jsonify({"msg": "已确认，后续检测将排除此 pair"})


@qc_bp.route("/api/quality-check/filter-options", methods=["GET"])
@jwt_required()
def filter_options():
    """返回当前方案下质检问题的可筛选维度"""
    plan_id = request.args.get("plan_id", type=int)

    q = QualityIssue.query
    if plan_id:
        q = q.filter(QualityIssue.plan_id == plan_id)

    issue_ids = [i[0] for i in q.with_entities(QualityIssue.id).all()]
    if not issue_ids:
        return jsonify({"data": {"issue_types": [], "assessor_units": [], "assessed_units": []}})

    types = (
        db.session.query(QualityIssue.issue_type)
        .filter(QualityIssue.id.in_(issue_ids))
        .distinct().order_by(QualityIssue.issue_type).all()
    )
    issue_types = [t[0] for t in types if t[0]]

    assessors = (
        db.session.query(Unit.id, Unit.name)
        .join(Task, Task.id == QualityIssue.task_id)
        .filter(QualityIssue.id.in_(issue_ids))
        .distinct().order_by(Unit.name).all()
    )
    assessor_units = [{"id": a[0], "name": a[1]} for a in assessors]

    assessed = (
        db.session.query(Unit.id, Unit.name)
        .join(Task, Task.id == QualityIssue.task_id)
        .filter(QualityIssue.id.in_(issue_ids))
        .distinct().order_by(Unit.name).all()
    )
    assessed_units = [{"id": u[0], "name": u[1]} for u in assessed]

    return jsonify({"data": {"issue_types": issue_types, "assessor_units": assessor_units, "assessed_units": assessed_units}})


@qc_bp.route("/api/quality-check/issues/managed", methods=["GET"])
@jwt_required()
def list_managed_issues():
    """问题管理：分类返回已确认和已修复的问题"""
    plan_id = request.args.get("plan_id", type=int)
    page = request.args.get("page", 1, type=int)
    page_size = request.args.get("page_size", 100, type=int)
    page_size = min(page_size, 500)

    q = QualityIssue.query.options(
        joinedload(QualityIssue.task).joinedload(Task.unit),
        joinedload(QualityIssue.task).joinedload(Task.assessor_unit),
    ).filter(QualityIssue.status.in_(["confirmed", "resolved"]))
    if plan_id:
        q = q.filter(QualityIssue.plan_id == plan_id)

    total = q.count()
    items = q.order_by(QualityIssue.id.desc()).offset((page - 1) * page_size).limit(page_size).all()
    return jsonify({"data": {
        "items": [_issue_to_dict(i) for i in items],
        "confirmed_count": QualityIssue.query.filter_by(plan_id=plan_id, status="confirmed").count() if plan_id else 0,
        "resolved_count": QualityIssue.query.filter_by(plan_id=plan_id, status="resolved").count() if plan_id else 0,
        "total": total, "page": page, "page_size": page_size,
    }})


@qc_bp.route("/api/quality-check/issues/grouped", methods=["GET"])
@jwt_required()
def list_grouped_issues():
    """问题管理：按主要任务合并显示确认无误和已更正的问题"""
    plan_id = request.args.get("plan_id", type=int)

    q = QualityIssue.query.filter(QualityIssue.status.in_(["confirmed", "resolved"]))
    if plan_id:
        q = q.filter(QualityIssue.plan_id == plan_id)

    issues = q.all()

    task_ids = list(set(iss.task_id for iss in issues if iss.task_id))
    tasks_map = {}
    if task_ids:
        preloaded_tasks = Task.query.options(
            joinedload(Task.assessor_unit), joinedload(Task.unit),
        ).filter(Task.id.in_(task_ids)).all()
        tasks_map = {t.id: t for t in preloaded_tasks}

    groups = {}
    for iss in issues:
        task = tasks_map.get(iss.task_id) if iss.task_id else None
        mt = task.main_task if task else "未知"
        assessor_name = task.assessor_unit.name if task and task.assessor_unit else "未知"
        unit_name = task.unit.name if task and task.unit else "未知"
        group_key = (mt, assessor_name)

        if group_key not in groups:
            groups[group_key] = {
                "main_task": mt, "assessor_name": assessor_name,
                "confirmed": [], "resolved": [], "units": set(),
            }
        groups[group_key]["units"].add(unit_name)
        if iss.status == "confirmed":
            groups[group_key]["confirmed"].append(_issue_to_dict(iss))
        else:
            groups[group_key]["resolved"].append(_issue_to_dict(iss))

    result = []
    for (mt, an), g in groups.items():
        result.append({
            "main_task": g["main_task"], "assessor_name": g["assessor_name"],
            "unit_names": sorted(g["units"]), "unit_count": len(g["units"]),
            "confirmed_count": len(g["confirmed"]), "resolved_count": len(g["resolved"]),
            "confirmed_items": g["confirmed"][:20], "resolved_items": g["resolved"][:20],
            "confirmed_total": len(g["confirmed"]), "resolved_total": len(g["resolved"]),
        })

    result.sort(key=lambda x: -(x["confirmed_count"] + x["resolved_count"]))
    return jsonify({"data": result})


# ==================== 代理指标 API ====================


def _proxy_pair_to_dict(pr):
    """将 ProxyMetricPair 转为前端需要的 dict"""
    task_a = pr.task_a
    task_b = pr.task_b
    middle_unit = pr.middle_unit

    middle_unit_name = middle_unit.name if middle_unit else ""
    source_unit_name = task_a.assessor_unit.name if task_a and task_a.assessor_unit else ""
    target_unit_name = task_b.unit.name if task_b and task_b.unit else ""

    recv_kw = (task_a.key_work or "") if task_a else ""
    recv_mt = (task_a.main_task or "") if task_a else ""
    asgn_kw = (task_b.key_work or "") if task_b else ""
    asgn_mt = (task_b.main_task or "") if task_b else ""

    return {
        "id": pr.id,
        "plan_id": pr.plan_id,
        "task_id_a": pr.task_id_a,
        "task_id_b": pr.task_id_b,
        "middle_unit_id": pr.middle_unit_id,
        "middle_unit_name": middle_unit_name,
        "source_unit_name": source_unit_name,
        "target_unit_name": target_unit_name,
        "task_a_summary": f"【{recv_kw[:50]}】{recv_mt[:50]}",
        "task_b_summary": f"【{asgn_kw[:50]}】{asgn_mt[:50]}",
        "recv_kw": recv_kw,
        "recv_mt": recv_mt,
        "asgn_kw": asgn_kw,
        "asgn_mt": asgn_mt,
        "similarity": pr.similarity,
        "confidence": pr.confidence,
        "status": pr.status,
        "remark": pr.remark or "",
        "created_at": pr.created_at.isoformat() if pr.created_at else "",
        "updated_at": pr.updated_at.isoformat() if pr.updated_at else "",
    }


@qc_bp.route("/api/proxy-metrics/pairs", methods=["GET"])
@jwt_required()
def list_proxy_pairs():
    """查询代理指标对列表"""
    plan_id = request.args.get("plan_id", type=int)
    status = request.args.get("status", "")
    middle_unit_id = request.args.get("middle_unit_id", type=int)
    search = request.args.get("search", "")
    page = request.args.get("page", 1, type=int)
    page_size = request.args.get("page_size", 50, type=int)
    page_size = min(page_size, 500)

    q = ProxyMetricPair.query.options(
        joinedload(ProxyMetricPair.task_a).joinedload(Task.assessor_unit),
        joinedload(ProxyMetricPair.task_a).joinedload(Task.unit),
        joinedload(ProxyMetricPair.task_b).joinedload(Task.assessor_unit),
        joinedload(ProxyMetricPair.task_b).joinedload(Task.unit),
        joinedload(ProxyMetricPair.middle_unit),
    )
    if plan_id:
        q = q.filter(ProxyMetricPair.plan_id == plan_id)
    if status:
        q = q.filter(ProxyMetricPair.status == status)
    if middle_unit_id:
        q = q.filter(ProxyMetricPair.middle_unit_id == middle_unit_id)
    if search:
        q = q.join(ProxyMetricPair.task_a).join(ProxyMetricPair.task_b).filter(
            or_(
                Task.key_work.contains(search),
                Task.main_task.contains(search),
            )
        )

    total = q.count()
    items = q.order_by(ProxyMetricPair.similarity.desc()).offset(
        (page - 1) * page_size
    ).limit(page_size).all()

    return jsonify({"data": {
        "items": [_proxy_pair_to_dict(pr) for pr in items],
        "total": total, "page": page, "page_size": page_size,
    }})


@qc_bp.route("/api/proxy-metrics/pairs/<int:pair_id>", methods=["GET"])
@jwt_required()
def get_proxy_pair(pair_id):
    """获取单个代理指标对详情"""
    pr = ProxyMetricPair.query.options(
        joinedload(ProxyMetricPair.task_a).joinedload(Task.assessor_unit),
        joinedload(ProxyMetricPair.task_a).joinedload(Task.unit),
        joinedload(ProxyMetricPair.task_b).joinedload(Task.assessor_unit),
        joinedload(ProxyMetricPair.task_b).joinedload(Task.unit),
        joinedload(ProxyMetricPair.middle_unit),
    ).get(pair_id)
    if not pr:
        return jsonify({"msg": "记录不存在"}), 404
    return jsonify({"data": _proxy_pair_to_dict(pr)})


@qc_bp.route("/api/proxy-metrics/pairs/<int:pair_id>", methods=["PUT"])
@jwt_required()
@require_perm("quality_manage")
def update_proxy_pair(pair_id):
    """更新代理指标对（状态/备注）"""
    pr = ProxyMetricPair.query.get(pair_id)
    if not pr:
        return jsonify({"msg": "记录不存在"}), 404
    data = request.get_json() or {}
    if "status" in data:
        pr.status = data["status"]
    if "remark" in data:
        pr.remark = data["remark"]
    pr.updated_at = datetime.utcnow()
    db.session.commit()
    return jsonify({"msg": "更新成功"})


@qc_bp.route("/api/proxy-metrics/pairs/batch-confirm", methods=["POST"])
@jwt_required()
@require_perm("quality_manage")
def batch_confirm_proxy_pairs():
    """批量确认代理指标对为误报"""
    data = request.get_json() or {}
    ids = data.get("ids", [])
    if not ids:
        return jsonify({"msg": "请选择要确认的记录"}), 400
    updated = ProxyMetricPair.query.filter(
        ProxyMetricPair.id.in_(ids)
    ).update(
        {"status": "confirmed", "updated_at": db.func.now()},
        synchronize_session=False,
    )
    db.session.commit()
    return jsonify({"msg": f"已确认 {updated} 条记录"})


@qc_bp.route("/api/proxy-metrics/pairs/batch-resolve", methods=["POST"])
@jwt_required()
@require_perm("quality_manage")
def batch_resolve_proxy_pairs():
    """批量标记代理指标对为已修正"""
    data = request.get_json() or {}
    ids = data.get("ids", [])
    if not ids:
        return jsonify({"msg": "请选择要标记的记录"}), 400
    updated = ProxyMetricPair.query.filter(
        ProxyMetricPair.id.in_(ids)
    ).update(
        {"status": "resolved", "updated_at": db.func.now()},
        synchronize_session=False,
    )
    db.session.commit()
    return jsonify({"msg": f"已标记 {updated} 条记录为已修正"})


@qc_bp.route("/api/proxy-metrics/pairs/<int:pair_id>/remark", methods=["PUT"])
@jwt_required()
@require_perm("quality_manage")
def update_proxy_pair_remark(pair_id):
    """更新代理指标对备注"""
    pr = ProxyMetricPair.query.get(pair_id)
    if not pr:
        return jsonify({"msg": "记录不存在"}), 404
    data = request.get_json() or {}
    pr.remark = data.get("remark", "")
    pr.updated_at = datetime.utcnow()
    db.session.commit()
    return jsonify({"msg": "备注已更新"})


@qc_bp.route("/api/proxy-metrics/filter-options", methods=["GET"])
@jwt_required()
def proxy_filter_options():
    """返回代理指标可筛选维度"""
    plan_id = request.args.get("plan_id", type=int)

    q = ProxyMetricPair.query
    if plan_id:
        q = q.filter(ProxyMetricPair.plan_id == plan_id)

    pair_ids = [r[0] for r in q.with_entities(ProxyMetricPair.id).all()]
    if not pair_ids:
        return jsonify({"data": {
            "middle_units": [],
            "status_counts": {"pending": 0, "confirmed": 0, "resolved": 0, "ignored": 0},
        }})

    # 中间单位列表
    middle_units = (
        db.session.query(Unit.id, Unit.name)
        .join(ProxyMetricPair, ProxyMetricPair.middle_unit_id == Unit.id)
        .filter(ProxyMetricPair.id.in_(pair_ids))
        .distinct().order_by(Unit.name).all()
    )

    # 各状态数量
    status_counts = {}
    for s in ["pending", "confirmed", "resolved", "ignored"]:
        status_counts[s] = ProxyMetricPair.query.filter(
            ProxyMetricPair.id.in_(pair_ids),
            ProxyMetricPair.status == s,
        ).count()

    return jsonify({"data": {
        "middle_units": [{"id": u[0], "name": u[1]} for u in middle_units],
        "status_counts": status_counts,
    }})


@qc_bp.route("/api/proxy-metrics/export", methods=["GET"])
@jwt_required()
def export_proxy_metrics():
    """导出代理指标清单"""
    plan_id = request.args.get("plan_id", type=int)
    status = request.args.get("status", "")

    q = ProxyMetricPair.query.options(
        joinedload(ProxyMetricPair.task_a).joinedload(Task.assessor_unit),
        joinedload(ProxyMetricPair.task_a).joinedload(Task.unit),
        joinedload(ProxyMetricPair.task_b).joinedload(Task.assessor_unit),
        joinedload(ProxyMetricPair.task_b).joinedload(Task.unit),
        joinedload(ProxyMetricPair.middle_unit),
    )
    if plan_id:
        q = q.filter(ProxyMetricPair.plan_id == plan_id)
    if status:
        q = q.filter(ProxyMetricPair.status == status)

    pairs = q.order_by(ProxyMetricPair.similarity.desc()).all()

    if not Workbook:
        return jsonify({"msg": "openpyxl 未安装"}), 500

    wb = Workbook()
    ws = wb.active
    ws.title = "疑似以指标考指标"

    headers = ["序号", "中间单位", "源主考单位", "目标被考核单位",
               "接收方重点工作", "接收方主要任务",
               "下发方重点工作", "下发方主要任务",
               "相似度", "置信度", "状态", "备注"]
    ws.append(headers)

    status_label = {"pending": "待处理", "confirmed": "已确认", "resolved": "已修正", "ignored": "已忽略"}
    for idx, pr in enumerate(pairs, 1):
        d = _proxy_pair_to_dict(pr)
        ws.append([
            idx,
            d["middle_unit_name"], d["source_unit_name"], d["target_unit_name"],
            d["recv_kw"], d["recv_mt"],
            d["asgn_kw"], d["asgn_mt"],
            f"{d['similarity']:.2%}" if d["similarity"] else "",
            d["confidence"],
            status_label.get(d["status"], d["status"]),
            d["remark"],
        ])

    col_widths = [6, 16, 16, 16, 30, 40, 30, 40, 10, 10, 10, 20]
    for col_idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="疑似以指标考指标清单.xlsx",
    )


@qc_bp.route("/api/quality-check/export", methods=["GET"])
@jwt_required()
def export_issues():
    """导出质检问题清单（全量或筛选后）"""
    plan_id = request.args.get("plan_id", type=int)
    mode = request.args.get("mode", "all")
    status = request.args.get("status", "")
    issue_type = request.args.get("issue_type", "")
    assessor_unit_id = request.args.get("assessor_unit_id", type=int)
    unit_id = request.args.get("unit_id", type=int)
    search = request.args.get("search", "")

    q = QualityIssue.query.options(
        joinedload(QualityIssue.task).joinedload(Task.unit),
        joinedload(QualityIssue.task).joinedload(Task.assessor_unit),
    )
    if plan_id:
        q = q.filter(QualityIssue.plan_id == plan_id)

    if mode == "filtered":
        if status:
            q = q.filter(QualityIssue.status == status)
        if issue_type:
            q = q.filter(QualityIssue.issue_type == issue_type)
        if assessor_unit_id:
            q = q.join(QualityIssue.task).filter(Task.assessor_unit_id == assessor_unit_id)
        if unit_id:
            q = q.join(QualityIssue.task).filter(Task.unit_id == unit_id)
        if search:
            q = q.filter(
                or_(
                    QualityIssue.text.contains(search),
                    QualityIssue.suggestion.contains(search),
                    QualityIssue.context.contains(search),
                )
            )

    issues = q.order_by(QualityIssue.id).all()

    if not Workbook:
        return jsonify({"msg": "openpyxl 未安装"}), 500

    wb = Workbook()
    ws = wb.active
    ws.title = "质检问题清单"

    headers = ["序号", "问题类型", "主考单位", "被考核单位", "主要任务", "问题描述", "修改建议", "状态", "置信度"]
    ws.append(headers)

    for idx, i in enumerate(issues, 1):
        unit_name = i.task.unit.name if i.task and i.task.unit else ""
        assessor_name = i.task.assessor_unit.name if i.task and i.task.assessor_unit else ""
        main_task = i.task.main_task if i.task else ""
        status_label = {"pending": "待处理", "confirmed": "已研判", "resolved": "已修复"}.get(i.status, i.status or "")
        ws.append([
            idx, i.issue_type or "", assessor_name, unit_name, main_task,
            i.text or "", i.suggestion or "", status_label, i.confidence or "",
        ])

    col_widths = [6, 18, 16, 16, 36, 50, 50, 10, 10]
    for col_idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    mode_label = "全量" if mode == "all" else "筛选结果"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"质检问题清单_{mode_label}.xlsx",
    )
